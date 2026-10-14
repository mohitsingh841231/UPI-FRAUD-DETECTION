from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(r"D:\UPI")
TEMPLATE=Path(r"C:\Users\mohit\Desktop\Fingerprint_Interim_Project_Report_I_Final_Tables (1).docx")
OUT=ROOT/"UPI_Fraud_Detection_Interim_Project_Report.docx"
FIG=ROOT/"tmp"/"interim_report_figures"; FIG.mkdir(parents=True,exist_ok=True)
PAPER_FIG=ROOT/"tmp"/"paper_figures"
NAVY=(28,55,82); BLUE=(47,111,154); LIGHT=(230,238,243); INK=(35,42,48); MUTED=(98,110,119); GRID=(210,218,223)
FONT=r"C:\Windows\Fonts\arial.ttf"; FONTB=r"C:\Windows\Fonts\arialbd.ttf"
def ft(n,b=False):
    try:return ImageFont.truetype(FONTB if b else FONT,n)
    except OSError:return ImageFont.load_default()

def make_workflow(path):
    im=Image.new('RGB',(1800,480),'white'); d=ImageDraw.Draw(im)
    title='UPI fraud detection project workflow'; d.text((900,35),title,font=ft(38,True),fill=NAVY,anchor='ma')
    labels=[('250,000-row','transaction CSV'),('Remove direct','identifiers'),('Encode and impute','feature pipeline'),('Stratified split','80% / 20%'),('Random forest','100 trees'),('Test metrics and','saved model'),('Flask','website and API')]
    xs=[22,275,528,781,1034,1287,1540]; y=170; w=220; h=150
    for i,(a,b) in enumerate(labels):
        x=xs[i]; d.rounded_rectangle((x,y,x+w,y+h),radius=18,fill=LIGHT,outline=BLUE,width=4)
        d.text((x+w/2,y+47),a,font=ft(24,True),fill=NAVY,anchor='mm')
        d.text((x+w/2,y+91),b,font=ft(23),fill=INK,anchor='mm')
        if i<len(xs)-1:
            sx=x+w+4; ex=xs[i+1]-9; cy=y+h/2
            d.line((sx,cy,ex,cy),fill=BLUE,width=6)
            d.polygon([(ex,cy),(ex-16,cy-11),(ex-16,cy+11)],fill=BLUE)
    d.text((900,390),'IDs, timestamps, and UPI addresses are excluded from model features.',font=ft(22),fill=MUTED,anchor='mm')
    im.save(path,dpi=(300,300))

make_workflow(FIG/'workflow.png')

# Keep a true copy of the supplied report package as the design authority.
doc=Document(TEMPLATE)
body=doc._element.body
logo=next((e for e in body.iterchildren() if e.tag==qn('w:p') and e.xpath('.//a:blip')),None)
sectPr=body.sectPr
for el in list(body.iterchildren()):
    if el is logo or el is sectPr: continue
    body.remove(el)
if logo is not None: body.remove(logo)
if logo is not None: body.insert(0,logo)

# The source report has one US Letter portrait section. Preserve it and the source theme/styles.
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=Inches(.9); sec.right_margin=Inches(.9); sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.75)
footer=sec.footer.paragraphs[0]
if footer.runs: footer.runs[0].text="Interim Project Report-I  |  UPI Fraud Detection Using Machine Learning"
else: footer.add_run("Interim Project Report-I  |  UPI Fraud Detection Using Machine Learning")
for r in footer.runs[1:]: r.text=""

def set_run(run,size=None,bold=None,italic=None,color=None,name='Times New Roman'):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    if size is not None: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic
    if color is not None: run.font.color.rgb=RGBColor(*color)
    return run

def addp(text='',style='Normal',align=None,before=0,after=5,keep=False,first=None,size=None,bold=None,italic=None):
    p=doc.add_paragraph(style=style)
    if text:
        r=p.add_run(text)
        if size is not None or bold is not None or italic is not None: set_run(r,size,bold,italic)
    if align is not None:p.alignment=align
    p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.keep_with_next=keep
    if first is not None:p.paragraph_format.first_line_indent=Inches(first)
    return p

def pagebreak():
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.add_run().add_break(WD_BREAK.PAGE)

def heading(text,level=1):
    p=addp(text,f'Heading {level}',after=5 if level==1 else 3,before=6 if level==1 else 3,keep=True)
    return p

def para(text):
    return addp(text,'Normal',after=5,first=.28)

def caption(text):
    p=addp(text,'Normal',align=WD_ALIGN_PARAGRAPH.CENTER,before=2,after=4,keep=True)
    for r in p.runs:set_run(r,9,False,False,(40,48,54))
    return p

def fig(path,width=6.1):
    p=addp('',align=WD_ALIGN_PARAGRAPH.CENTER,after=0,keep=True)
    p.add_run().add_picture(str(path),width=Inches(width))
    return p

def fill_cell(cell,text,bold=False,center=False,size=9):
    cell.text=''; p=cell.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
    r=p.add_run(str(text));set_run(r,size,bold)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade(cell,fill='E5EBEF'):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),fill);tcPr.append(sh)

def table(rows,widths=None,font_size=9):
    cols=len(rows[0]); t=doc.add_table(rows=len(rows),cols=cols);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Inches(w)
    for i,row in enumerate(rows):
        for j,val in enumerate(row):
            fill_cell(t.cell(i,j),val,bold=(i==0),center=(i==0 or j>0),size=font_size)
            if i==0:shade(t.cell(i,j))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def titled_table(label,title,rows,widths=None,font_size=9):
    p=addp(label,align=WD_ALIGN_PARAGRAPH.CENTER,before=5,after=1,keep=True)
    for r in p.runs:set_run(r,10,True)
    p=addp(title,align=WD_ALIGN_PARAGRAPH.CENTER,after=3,keep=True)
    for r in p.runs:set_run(r,9,True)
    return table(rows,widths,font_size)

# Cover page. The original institutional mark is retained from the source template.
if logo is not None:
    try:
        from docx.text.paragraph import Paragraph
        Paragraph(logo,doc).alignment=WD_ALIGN_PARAGRAPH.CENTER
    except Exception: pass
for _ in range(2): addp('',after=8)
for text,size,bold,after in [
    ('Interim Project Report-I',19,True,2),('on',12,True,4),
    ('UPI Fraud Detection Using Machine Learning',20,True,20),
    ('Submitted in partial fulfillment for the award of',12,False,3),
    ('BACHELOR OF TECHNOLOGY',15,True,4),('Degree in',12,False,2),
    ('COMPUTER SCIENCE & ENGINEERING',14,True,5),('2026–27',12,True,20),
    ('Under the Guidance of: [GUIDE NAME]',11,False,3),
    ('Submitted By: Mohit Singh  |  Admission Number: [NOT PROVIDED]',11,False,18),
    ('SCHOOL OF COMPUTER SCIENCE & ENGINEERING',12,True,3),
    ('GALGOTIAS UNIVERSITY',15,True,4),('OCTOBER 2026',12,True,0)]:
    p=addp(text,align=WD_ALIGN_PARAGRAPH.CENTER,after=after,keep=True)
    for r in p.runs:set_run(r,size,bold)
pagebreak()

# Project progress form
heading('PROJECT PROGRESS REPORT MTE / MID TERM EXAM',1)
addp('B.Tech. Project Details — Semester: 7th',align=WD_ALIGN_PARAGRAPH.CENTER,after=6)
titled_table('TABLE 1.1','Project Details',[
 ['Field','Details'],['Project ID','[NOT PROVIDED]'],['Project Title','UPI Fraud Detection Using Machine Learning'],
 ['Project Progress','Dataset reviewed; preprocessing and random-forest pipeline trained and evaluated; Jupyter notebook and Flask demonstration website prepared.'],
 ['Research Paper Title','UPI Fraud Detection Using Machine Learning'],['Research Report Progress','Interim project report prepared from the implemented project and measured test results.'],
 ['Suggestions by Guide','[TO BE COMPLETED BY GUIDE]'],['Suggestions by Examiner','[TO BE COMPLETED BY EXAMINER]']], [2.0,4.9],8)
addp('Student Progress Details (Filled by Guide Only)',align=WD_ALIGN_PARAGRAPH.CENTER,before=5,after=3,bold=True)
table([['S. No.','Student Name','Admission Number','Performance','Approval'],
       ['1','Mohit Singh','[NOT PROVIDED]','☐ Excellent  ☐ Good  ☐ Satisfactory','☐ Approved  ☐ Not Approved'],
       ['2','[TEAM MEMBER, IF ANY]','[ADMISSION NO.]','','']], [0.55,1.35,1.3,1.65,1.65],8)
addp('Supervisor Signature: __________________________       Examiner Signature: __________________________',align=WD_ALIGN_PARAGRAPH.CENTER,before=5,after=0)
pagebreak()

# Synopsis
heading('SYNOPSIS',1)
addp('Project Title: UPI Fraud Detection Using Machine Learning',bold=True,after=5)
para('Unified Payments Interface (UPI) supports rapid person-to-person and merchant payments. This project studies a supervised machine-learning pipeline that classifies transaction records as fraud or legitimate and provides a Flask demonstration interface for displaying evaluation metrics and scoring an example transaction.')
para('Problem statement. The project asks whether transaction attributes in the supplied CSV can support binary fraud classification on a held-out sample. The available file contains 250,000 rows and the project does not document its collection method, label rules, or risk-score generation process; results therefore describe this dataset rather than live UPI traffic.')
para('Objectives. Inspect the transaction data; exclude direct identifiers; preprocess numeric and categorical fields; train a random-forest classifier; report accuracy, precision, recall, F1, ROC AUC and the confusion matrix; and connect the saved model to a local Flask website.')
para('Methodology. The pipeline uses a reproducible stratified 80:20 train-test split, median and most-frequent imputation, one-hot encoding, and a 100-tree random forest with class weighting. Transaction ID, timestamp and UPI addresses are omitted as predictors. A notebook documents the evaluation, while Flask provides a dashboard and JSON prediction endpoints.')
para('Results and limitations. On the 50,000-row test set, accuracy was 92.98%, precision 93.09%, recall 92.85%, F1 92.97%, and ROC AUC 0.930. The test confusion matrix was [[23,275, 1,722], [1,787, 23,216]]. The file has a 50.006% fraud rate. The unknown provenance and the strong contribution of `risk_score` require follow-up checks for target leakage and external validity.')
addp('Keywords: UPI, transaction fraud, random forest, confusion matrix, ROC AUC, Flask.',after=0,italic=True)
pagebreak()

# Table of contents and report lists
heading('TABLE OF CONTENTS',1)
table([['CONTENTS','PAGE'],['Synopsis',''],['List of Figures',''],['List of Tables',''],['List of Symbols and Abbreviations',''],
       ['Chapter 1: Introduction',''],['Chapter 2: Literature Review',''],['Chapter 3: Proposed Methodology',''],
       ['Chapter 4: Software Implementation',''],['Chapter 5: Results, Discussion and Future Work',''],['Chapter 6: Conclusion',''],['References','']], [5.8,1.0],9)
addp('Page numbers can be updated in Word after final pagination.',italic=True,after=2)
heading('LIST OF FIGURES',1)
table([['FIGURE','DESCRIPTION','PAGE'],['Figure 3.1','UPI fraud detection project workflow',''],['Figure 5.1','Confusion matrix on the held-out test set',''],['Figure 5.2','ROC curve for held-out predictions',''],['Figure 5.3','Random-forest feature importance','']], [1.0,4.9,.9],8)
heading('LIST OF TABLES',1)
table([['TABLE','DESCRIPTION','PAGE'],['Table 1.1','Project details and progress',''],['Table 1.2','Project objectives and deliverables',''],['Table 2.1','Related approaches and relevance',''],['Table 2.2','Research gaps and project response',''],['Table 3.1','Model configuration',''],['Table 4.1','Software tools and roles',''],['Table 5.1','Test-set evaluation metrics',''],['Table 5.2','Confusion-matrix counts','']], [1.0,4.9,.9],8)
heading('LIST OF SYMBOLS AND ABBREVIATIONS',1)
table([['TERM','MEANING'],['UPI','Unified Payments Interface'],['P2P','Person-to-person transaction'],['P2M','Person-to-merchant transaction'],['RF','Random forest'],['ROC','Receiver operating characteristic'],['AUC','Area under the ROC curve'],['TP / TN','True positive / true negative'],['FP / FN','False positive / false negative'],['API','Application programming interface']], [1.4,5.4],8)
pagebreak()

# Chapter 1
heading('CHAPTER 1: INTRODUCTION',1)
para('This chapter sets out the project context, research question, objectives and scope. It also states the limits of the current evidence so the reported model scores are not mistaken for operational fraud-detection performance.')
heading('1.1 Background of the Problem',2)
para('UPI is an instant payment system that enables bank-account transfers through interoperable applications [5]. High transaction speed makes automated risk screening useful as one part of a broader review process. A classifier can learn associations in past labelled transactions, but its value depends on the quality and timing of the input features and the reliability of the labels.')
para('The supplied dataset contains transaction amount, transaction type, status, sender and receiver age groups, state, banks, device and network, time-of-day fields, a risk score, and a binary fraud label. The project trains a random forest to predict that label and exposes the saved model through a small website.')
heading('1.2 Motivation and Significance',2)
para('The project brings together the main stages of a practical machine-learning exercise: inspecting a CSV, preparing mixed data types, separating training from testing, evaluating predictions, saving a model, and connecting the model to an application. Showing the confusion matrix alongside accuracy helps distinguish false alerts from missed fraud cases.')
heading('1.3 Problem Statement',2)
para('Given the features in the provided transaction file, can a supervised classifier distinguish rows labelled fraud from rows labelled legitimate on a held-out test split, and can the trained pipeline return an understandable score through a web interface?')
heading('1.4 Objectives of the Project',2)
for s in ['Review the supplied dataset and identify the target and available transaction features.','Remove direct identifiers from the predictor set and preprocess numeric and categorical fields.','Train a reproducible random-forest classifier using a stratified train-test split.','Evaluate the model with precision, recall, F1, ROC AUC, accuracy and a confusion matrix.','Present test results and a transaction-scoring form through a local Flask website.']:
    p=addp(s,after=3,first=.18);p.style='List Bullet'
titled_table('TABLE 1.2','Project Objectives and Deliverables',[
 ['NO.','OBJECTIVE','DELIVERABLE','STATUS'],
 ['1','Inspect transaction CSV','Dataset summary and target profile','Completed'],
 ['2','Train and evaluate classifier','Saved pipeline and held-out metrics','Completed'],
 ['3','Document results','Notebook with metric and matrix outputs','Prepared'],
 ['4','Serve model output','Flask dashboard and prediction endpoint','Prepared'],
 ['5','Validate real-world use','Provenance, temporal and leakage checks','Future work']], [.45,2.2,2.6,1.05],8)
heading('1.5 Scope of the Project',2)
para('This report covers the supplied CSV, the existing scikit-learn pipeline, its single 80:20 stratified evaluation, the confusion matrix and feature importances, and a local Flask demonstration. It does not verify the dataset source, establish label adjudication, compare multiple algorithms, or test a live UPI feed.')
heading('1.6 Limitations of the Study',2)
para('The dataset has an almost 50:50 target distribution, but no provenance or label-generation documentation is present in the project. A row-wise random split may place related transactions in both partitions. The risk-score field is strongly associated with the target and may contain information unavailable at decision time. These issues limit conclusions about generalization.')
heading('1.7 Organization of the Report',2)
para('Chapter 2 reviews relevant approaches and research gaps. Chapter 3 documents the dataset, feature treatment, classifier and evaluation design. Chapter 4 describes the Python and Flask implementation. Chapter 5 presents the measured results, figures and limitations. Chapter 6 concludes the report and identifies follow-up work.')
pagebreak()

# Chapter 2
heading('CHAPTER 2: LITERATURE REVIEW',1)
para('Research on UPI fraud detection motivates using transaction records to identify suspicious patterns. The supplied chapter surveys machine-learning approaches and discusses feature engineering and monitoring [1]. It is useful as background, but it does not establish the origin or validity of the dataset used in this project. The two PDFs supplied with the request are identical copies of that same chapter and are cited once.')
heading('2.1 Review of Existing Systems and Techniques',2)
heading('Rule-based screening',3)
para('Rules flag transactions that meet predefined conditions, such as limits or known risk indicators. Rules are easy to inspect, but require ongoing maintenance and may not capture interactions among multiple transaction fields.')
heading('Supervised classification',3)
para('Supervised models learn a mapping from historical feature rows to supplied labels. Random forests combine tree predictors and can represent nonlinear relationships in mixed feature sets [2]. Their scores remain tied to the data and labels used during training.')
heading('Anomaly detection and monitoring',3)
para('Anomaly methods look for observations that differ from a learned baseline, while monitoring tracks changes after deployment. The supplied UPI chapter discusses these directions [1], but the present project implements a static supervised classifier and does not evaluate online adaptation or drift detection.')
heading('2.2 Review of Related Research Work',2)
para('Jagadeesan et al. review machine-learning concepts for UPI fraud detection and describe the need to consider feature construction and changing fraud patterns [1]. Breiman’s random-forest formulation provides the basis for the ensemble classifier used here [2]. Fawcett explains ROC curves as a way to examine classifier behaviour across thresholds and cautions against relying on a single accuracy value [3]. Pedregosa et al. describe the scikit-learn toolkit used to assemble and evaluate this pipeline [4].')
heading('2.3 Comparative Analysis of Existing Approaches',2)
titled_table('TABLE 2.1','Related Approaches and Their Relevance',[
 ['APPROACH','STRENGTH','LIMITATION','RELEVANCE'],
 ['Rule-based screening','Transparent conditions','Needs manual updates; limited interactions','Useful operational baseline, not implemented here'],
 ['Random forest','Ensemble of trees; handles nonlinear patterns','Importance and scores depend on data quality','Classifier evaluated in this project'],
 ['Anomaly detection','Can identify patterns outside a baseline','Needs a defined baseline and review process','Discussed in related work; not evaluated here'],
 ['Flask demonstration','Shows saved metrics and one-row scoring','No authentication, monitoring or payment integration','Prototype interface for this project']], [1.25,1.65,2.05,1.35],8)
heading('2.4 Identification of Research Gaps',2)
titled_table('TABLE 2.2','Research Gaps and Project Response',[
 ['NO.','GAP OR UNKNOWN','PROJECT RESPONSE / NEXT REQUIREMENT'],
 ['1','Dataset source and label rules are undocumented','Obtain provenance and label-adjudication details'],
 ['2','Fraud prevalence is 50.006% in the file','Evaluate at a representative operational prevalence'],
 ['3','Risk-score timing and derivation are unknown','Run an ablation without risk_score and verify availability at prediction time'],
 ['4','Random row split may share entities across partitions','Use time-based and entity-grouped evaluation'],
 ['5','No cost-based threshold study is reported','Set threshold using review capacity and error costs'],
 ['6','One model and one split were assessed','Add baselines, repeated validation and external testing']], [.45,2.6,3.25],8)
heading('2.5 Summary of Literature Findings',2)
para('Prior work supports testing supervised classifiers on transaction features, but the usefulness of a model must be judged on a clearly described dataset and evaluation design. This report emphasizes multiple metrics, false-positive and false-negative counts, and the limits imposed by missing data provenance. No claim of superiority over other classifiers is made because the project did not run a comparative experiment.')
pagebreak()

# Chapter 3
heading('CHAPTER 3: PROPOSED METHODOLOGY',1)
para('The implemented workflow reads the supplied CSV, separates the binary target, drops direct identifiers, fits preprocessing and a random forest on the training partition, and evaluates predictions on the held-out partition. The saved pipeline is then loaded by the Flask application.')
heading('3.1 System Overview',2)
addp('Figure 3.1. UPI fraud detection project workflow.',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,after=2)
fig(FIG/'workflow.png',6.15)
heading('3.2 Dataset Preparation',2)
para('The dataset has 250,000 records and 20 columns, including the target `is_fraud`. It contains 125,016 fraud labels and 124,984 legitimate labels, a fraud rate of 50.0064%. Each record includes transaction information such as amount, type and status; sender/receiver categories; device and network information; time features; risk_score; and the target. The project files do not identify the dataset publisher, collection process, or label-review protocol.')
heading('3.3 Feature Selection and Preprocessing',2)
para('The pipeline removes `transaction id`, `timestamp`, `sender_upi_id`, `receiver_upi_id`, and `is_fraud` from the predictor matrix. The remaining 15 features are split into numeric and categorical groups. Numeric missing values are filled with the training median and scaled. Categorical missing values are filled with the training mode and one-hot encoded with unseen categories ignored. These operations are part of the fitted scikit-learn pipeline so the test set does not determine preprocessing values.')
heading('3.4 Random-Forest Classifier',2)
para('The project uses a random forest with 100 trees, maximum depth 20, minimum leaf size 2, `balanced_subsample` class weights, random seed 42, and four worker threads. A random forest combines tree predictions to form a classification result [2]. The weighting option operates during tree fitting; it does not make the test data representative of real-world fraud prevalence.')
titled_table('TABLE 3.1','Model Configuration',[
 ['SETTING','VALUE','PURPOSE'],['Estimator','RandomForestClassifier','Binary classification'],['Trees','100','Number of tree predictors'],
 ['Maximum depth','20','Limits tree depth'],['Minimum leaf size','2','Minimum records per leaf'],['Class weighting','balanced_subsample','Weights bootstrap samples by class frequency'],
 ['Split','80% train / 20% test','50,000 records held out'],['Random seed','42','Reproducible partition and estimator'],['Target','is_fraud','1 = fraud; 0 = legitimate']], [1.55,1.85,2.9],8)
heading('3.5 Training and Test Protocol',2)
para('The data are divided with `train_test_split`, test_size 0.20, random_state 42, and stratification on the target. Training uses 200,000 rows and testing uses 50,000. Preprocessing and model fitting occur on the training rows only. The split preserves class proportions but is not time-based or grouped by participant, bank, or merchant.')
heading('3.6 Evaluation Metrics',2)
para('Accuracy measures the share of correct predictions. Precision is the share of predicted fraud rows that are labelled fraud; recall is the share of labelled fraud rows detected. F1 summarizes precision and recall with their harmonic mean. ROC AUC measures ranking across thresholds [3]. The confusion matrix gives counts for true negatives, false positives, false negatives, and true positives. The project’s web interface labels a transaction as potential fraud when its predicted probability is at least 0.5.')
pagebreak()

# Chapter 4
heading('CHAPTER 4: SOFTWARE IMPLEMENTATION',1)
para('The project is implemented in Python. The Jupyter notebook documents dataset review, training, evaluation and visualization. A command-line training script writes the serialized model bundle, and the Flask server loads that bundle to serve a dashboard and prediction endpoint.')
heading('4.1 Technologies and Tools',2)
titled_table('TABLE 4.1','Software Tools and Their Roles',[
 ['TOOL','ROLE IN THE PROJECT'],['Python','Data processing, model training and web server'],['pandas','CSV loading and feature preparation'],
 ['scikit-learn','Preprocessing pipeline, random forest, split and metrics'],['joblib','Persistence of the trained pipeline and test metrics'],
 ['Jupyter Notebook','Interactive review of training and evaluation'],['Flask','Dashboard and JSON endpoints'],['HTML, CSS and JavaScript','Browser dashboard and transaction scoring form']], [2.0,4.3],8)
heading('4.2 Training and Model Persistence',2)
para('The `train.py` entry point calls the project training function. It reads the CSV, builds the preprocessing and classifier pipeline, fits on the training partition, calculates held-out metrics, and saves the estimator, input feature list and metrics together in `app/fraud_model.joblib`. Saving the feature list allows the server to construct a row with the expected model inputs.')
heading('4.3 Flask Dashboard and Prediction Workflow',2)
para('The home page displays accuracy, precision, recall, ROC AUC and the confusion matrix from the saved evaluation. `GET /api/metrics` returns the same metrics as JSON. `POST /api/predict` accepts a JSON transaction record, arranges its fields in the model’s input schema, obtains the fraud probability, and returns that probability with the 0/1 label at a 0.5 threshold.')
para('The browser form collects transaction amount, type, merchant category, status, age groups, state, banks, device/network, risk score and time features. Unknown categories are handled by the one-hot encoder. The current application is a local demonstration: it has no login, audit trail, transaction-system integration, alert review queue, or production monitoring.')
heading('4.4 Privacy and Responsible Use',2)
para('Direct transaction and UPI identifiers are not model inputs. The interface should remain local or be placed behind appropriate access controls before any wider use. A model score should not independently authorize or reject a payment; the dataset and evaluation do not support that operational decision.')
pagebreak()

# Chapter 5
heading('CHAPTER 5: RESULTS, DISCUSSION AND FUTURE WORK',1)
para('The classifier was evaluated once on 50,000 stratified test records. At the project threshold of 0.5, it correctly classified 46,491 records. The estimates below reflect only this held-out subset and its near-balanced label distribution.')
heading('5.1 Test-Set Performance',2)
titled_table('TABLE 5.1','Held-Out Test Metrics',[
 ['METRIC','VALUE','INTERPRETATION'],['Accuracy','92.98%','46,491 of 50,000 labels correct'],['Precision','93.09%','Among fraud predictions, share labelled fraud'],
 ['Recall','92.85%','Among fraud labels, share detected'],['F1 score','92.97%','Harmonic mean of precision and recall'],['ROC AUC','0.9300','Threshold-independent ranking summary'],
 ['Test rows','50,000','24,997 legitimate; 25,003 fraud']], [1.25,1.05,4.0],8)
titled_table('TABLE 5.2','Confusion-Matrix Counts at Threshold 0.5',[
 ['ACTUAL CLASS','PREDICTED LEGITIMATE','PREDICTED FRAUD'],['Legitimate','23,275 (TN)','1,722 (FP)'],['Fraud','1,787 (FN)','23,216 (TP)']], [1.6,2.55,2.55],8)
addp('Figure 5.1. Confusion matrix on the held-out test set.',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,after=2)
fig(PAPER_FIG/'confusion_matrix.png',6.15)
heading('5.2 Discussion',2)
para('Precision and recall are similar on this test set, and the ROC curve indicates a ROC AUC of 0.930. The 1,722 false positives and 1,787 false negatives show that the model makes both types of error. In a live payment population, a lower fraud prevalence would generally reduce the fraction of alerts that are truly fraudulent even if sensitivity and specificity stayed similar. Therefore, accuracy and this test precision cannot be transferred directly to operational UPI traffic.')
addp('Figure 5.2. ROC curve for held-out predictions; plotted points are interpolated at evenly spaced false-positive rates.',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,after=2)
fig(PAPER_FIG/'roc_curve.png',6.15)
heading('5.3 Feature-Importance Review',2)
para('The model’s impurity-based feature importance assigns 48.5% to risk_score and 39.5% to transaction amount; merchant category contributes about 5.9%. Together, risk_score and amount account for about 88.0% of the total. The risk score has a mean of 1,163.06 for fraud-labelled rows and 295.39 for legitimate rows. This association is descriptive, not causal. Since the source and time of creation of risk_score are undocumented, it may be a valid early feature or may encode label-related information. An ablation that retrains without it is needed before interpreting model performance.')
addp('Figure 5.3. Largest aggregated random-forest feature importances.',align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,after=2)
fig(PAPER_FIG/'feature_importance.png',6.15)
heading('5.4 Limitations and Future Work',2)
para('The target is almost exactly balanced in the supplied file, and its label provenance is unknown. The random row split does not test performance on later transactions or unseen participants. Only one model, one split and one threshold were assessed. Impurity importance can favour some feature types and should not be read as an explanation of individual model decisions. The output probabilities have not been calibrated for another population.')
para('Next steps are to document data collection and label rules; verify that risk_score is available before the prediction point; compare the current model with an ablation that removes risk_score and with simple baselines; run time-based and entity-grouped tests; and choose a threshold using realistic fraud prevalence, review capacity and error costs. A broader evaluation should also check calibration, subgroup error rates, and performance drift. These experiments are outside the current project results.')
pagebreak()

# Chapter 6 and bibliography
heading('CHAPTER 6: CONCLUSION',1)
para('This interim report documents a complete UPI fraud classification prototype built from the supplied transaction CSV. The project uses a stratified 80:20 split and a class-weighted random forest, reaching 92.98% accuracy and 0.930 ROC AUC on 50,000 held-out rows. It also presents precision, recall, F1, the confusion matrix, a ROC curve and feature importance, and serves the saved model through a local Flask dashboard.')
para('The measured result is encouraging within this file, but dataset provenance and the meaning of the risk-score feature remain unverified. A 50.006% fraud prevalence and random row split also limit how the scores can be interpreted beyond the test sample. The next evaluation should verify labels and feature timing, remove risk_score as a sensitivity check, and test on later and entity-separated records before proposing any operational use.')
heading('REFERENCES',1)
refs=[
 '[1] S. Jagadeesan, K. S. Arjun, G. Dhanika, G. Karthikeyan, and K. Deepika, “UPI fraud detection using machine learning,” in Challenges in Information, Communication and Computing Technology, 2025, pp. 755–760, doi: 10.1201/9781003559085-130.',
 '[2] L. Breiman, “Random forests,” Machine Learning, vol. 45, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.',
 '[3] T. Fawcett, “An introduction to ROC analysis,” Pattern Recognition Letters, vol. 27, no. 8, pp. 861–874, 2006, doi: 10.1016/j.patrec.2005.10.010.',
 '[4] F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning Research, vol. 12, no. 85, pp. 2825–2830, 2011.',
 '[5] National Payments Corporation of India, “UPI frequently asked questions,” accessed Oct. 14, 2026. [Online]. Available: https://www.npci.org.in/what-we-do/upi/faqs.',
 '[6] UPI Fraud Detection project dataset, `archive/dataset.csv`, user-supplied file, 250,000 transaction records; source, collection period and label-generation protocol not stated in the project files.'
]
for r in refs:
    p=addp(r,after=5,first=0)
    for run in p.runs:set_run(run,10)

# Page-number field retained in the template footer. Keep original style and text, updating topic only.
doc.core_properties.title='Interim Project Report-I on UPI Fraud Detection Using Machine Learning'
doc.core_properties.author='Mohit Singh'
doc.core_properties.subject='Interim project report based on the implemented UPI fraud detection project'
doc.core_properties.keywords='UPI, fraud detection, random forest, interim project report'
doc.save(OUT)
print(OUT)
