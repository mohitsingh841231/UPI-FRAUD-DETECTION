from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\UPI")
OUT = ROOT / "UPI_Fraud_Detection_Research_Paper.docx"
TMP = ROOT / "tmp" / "paper_figures"
TMP.mkdir(parents=True, exist_ok=True)

FONT = r"C:\Windows\Fonts\arial.ttf"
FONT_B = r"C:\Windows\Fonts\arialbd.ttf"
def font(size, bold=False):
    path = FONT_B if bold else FONT
    try: return ImageFont.truetype(path, size)
    except OSError: return ImageFont.load_default()

NAVY = (27, 48, 73); BLUE = (44, 105, 151); LIGHT = (224, 235, 241)
INK = (35, 42, 48); GRAY = (104, 115, 123); GRID = (213, 219, 223)

def draw_text(draw, xy, text, size=22, fill=INK, bold=False, anchor=None):
    draw.text(xy, text, font=font(size, bold), fill=fill, anchor=anchor)

def save_confusion(path):
    im = Image.new("RGB", (1500, 1050), "white"); d = ImageDraw.Draw(im)
    draw_text(d, (750, 50), "Confusion matrix on the held-out test set", 40, NAVY, True, "ma")
    x0,y0,cell=430,240,380
    draw_text(d,(x0+cell,170),"Predicted legitimate",25,GRAY,False,"ma")
    draw_text(d,(x0+2*cell,170),"Predicted fraud",25,GRAY,False,"ma")
    draw_text(d,(255,y0+cell//2),"Actual legitimate",25,GRAY,False,"mm")
    draw_text(d,(255,y0+3*cell//2),"Actual fraud",25,GRAY,False,"mm")
    vals=[[23275,1722],[1787,23216]]; mx=max(max(r) for r in vals)
    for r in range(2):
        for c in range(2):
            x=x0+c*cell; y=y0+r*cell; frac=vals[r][c]/mx
            color=(int(244-125*frac),int(248-82*frac),int(250-53*frac))
            d.rectangle((x,y,x+cell-12,y+cell-12),fill=color,outline=(255,255,255),width=3)
            draw_text(d,(x+(cell-12)//2,y+cell//2-18),f"{vals[r][c]:,}",46,NAVY,True,"mm")
            draw_text(d,(x+(cell-12)//2,y+cell//2+36),f"{vals[r][c]/50000:.1%} of test set",19,GRAY,False,"mm")
    draw_text(d,(750,1010),"Rows are actual labels; columns are model predictions.",20,GRAY,False,"ms")
    im.save(path,dpi=(300,300))

def save_roc(path):
    pts=[0.0012,0.6690,0.9308,0.9348,0.9382,0.9425,0.9466,0.9505,0.9543,0.9580,0.9617,0.9655,0.9694,0.9728,0.9757,0.9796,0.9841,0.9880,0.9923,0.9964,1.0]
    im=Image.new("RGB",(1450,1050),"white"); d=ImageDraw.Draw(im)
    draw_text(d,(725,45),"Receiver operating characteristic",40,NAVY,True,"ma")
    x0,y0,w,h=190,150,1080,735
    for i in range(6):
        x=x0+i*w/5; y=y0+h-i*h/5
        d.line((x,y0,x,y0+h),fill=GRID,width=2); d.line((x0,y,x0+w,y),fill=GRID,width=2)
        draw_text(d,(x,y0+h+25),f"{i/5:.1f}",20,GRAY,False,"ma")
        draw_text(d,(x0-25,y),f"{i/5:.1f}",20,GRAY,False,"rm")
    d.line((x0,y0+h,x0+w,y0),fill=(164,171,176),width=4)
    coords=[]
    for i,t in enumerate(pts):
        x=x0+i*w/20; y=y0+h-t*h; coords.append((int(x),int(y)))
    d.line(coords,fill=BLUE,width=9)
    for x,y in coords[::2]: d.ellipse((x-8,y-8,x+8,y+8),fill=BLUE)
    draw_text(d,(x0+w/2,y0+h+78),"False positive rate",23,GRAY,False,"ma")
    label=Image.new("RGBA",(300,50),(255,255,255,0)); ld=ImageDraw.Draw(label)
    ld.text((150,25),"True positive rate",font=font(23),fill=GRAY,anchor="mm")
    label=label.rotate(90,expand=True)
    im.paste(label,(28,y0+(h-label.height)//2),label)
    draw_text(d,(x0+w-210,y0+65),"AUC = 0.930",28,NAVY,True,"ma")
    im.save(path,dpi=(300,300))

def save_importance(path):
    items=[("risk_score",.4853),("amount (INR)",.3951),("merchant_category",.0586),("hour_of_day",.0099),("sender_state",.0077),("receiver_bank",.0069)]
    im=Image.new("RGB",(1500,850),"white"); d=ImageDraw.Draw(im)
    draw_text(d,(750,45),"Top input feature importances",40,NAVY,True,"ma")
    left,top,bw,row=390,170,840,96
    for i,(name,val) in enumerate(items):
        y=top+i*row
        draw_text(d,(left-22,y+25),name,25,INK,False,"rm")
        d.rounded_rectangle((left,y,left+bw,y+43),radius=8,fill=LIGHT)
        d.rounded_rectangle((left,y,left+int(bw*val),y+43),radius=8,fill=BLUE)
        draw_text(d,(left+bw+18,y+21),f"{val:.1%}",22,NAVY,True,"lm")
    draw_text(d,(750,790),"Impurity-based importances; association is not causation.",20,GRAY,False,"ms")
    im.save(path,dpi=(300,300))

figs={name:TMP/f"{name}.png" for name in ("confusion_matrix","roc_curve","feature_importance")}
save_confusion(figs["confusion_matrix"]); save_roc(figs["roc_curve"]); save_importance(figs["feature_importance"])

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)
sec.top_margin=Inches(.66); sec.bottom_margin=Inches(.66)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)

def set_font(run, size=9, bold=False, italic=False, color=INK, name="Times New Roman"):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:eastAsia'),name)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic
    run.font.color.rgb=RGBColor(*color)
    return run

normal=doc.styles['Normal']; normal.font.name="Times New Roman"; normal.font.size=Pt(9)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),"Times New Roman")
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after=Pt(0); normal.paragraph_format.space_before=Pt(0)
normal.paragraph_format.line_spacing=1.0

def para(text="",style=None,align=None,first=0,after=0,before=0,keep=False):
    p=doc.add_paragraph(style=style)
    if text: p.add_run(text)
    if align is not None: p.alignment=align
    p.paragraph_format.first_line_indent=Inches(first)
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.keep_with_next=keep
    return p

p=para(align=WD_ALIGN_PARAGRAPH.CENTER,after=8,before=26,keep=True)
set_font(p.add_run("UPI Fraud Detection with a Random Forest Model"),22,True,color=(0,0,0))
p=para(align=WD_ALIGN_PARAGRAPH.CENTER,after=2,keep=True)
set_font(p.add_run("Mohit Singh"),11,False,color=(0,0,0))
p=para(align=WD_ALIGN_PARAGRAPH.CENTER,after=0)
set_font(p.add_run("UPI Fraud Detection Project"),9,False,italic=True,color=GRAY)

sec2=doc.add_section(WD_SECTION.CONTINUOUS)
sec2.page_width=Inches(8.5); sec2.page_height=Inches(11)
sec2.left_margin=Inches(.68); sec2.right_margin=Inches(.68)
sec2.top_margin=Inches(.66); sec2.bottom_margin=Inches(.66)
cols=sec2._sectPr.xpath('./w:cols')[0]; cols.set(qn('w:num'),'2'); cols.set(qn('w:space'),'360')

def body(text, first=.14):
    p=para(first=first,after=2)
    set_font(p.add_run(text),9)
    return p

def section(title):
    p=para(align=WD_ALIGN_PARAGRAPH.CENTER,before=5,after=2,keep=True)
    set_font(p.add_run(title),9,True,color=(0,0,0))
    return p

def subhead(title):
    p=para(before=4,after=1,keep=True)
    set_font(p.add_run(title),9,True,italic=True,color=(0,0,0))
    return p

def table(rows,widths,header=True,font_size=8):
    t=doc.add_table(rows=0,cols=len(widths)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    tPr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'AAB4BA'); borders.append(el)
    tPr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].width=Inches(widths[j]); cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cells[j]._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
            for edge,v in [('top','50'),('start','65'),('bottom','50'),('end','65')]:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),v); e.set(qn('w:type'),'dxa'); mar.append(e)
            tcPr.append(mar)
            if header and ri==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5EBEF'); tcPr.append(sh)
            p=cells[j].paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (j>0 or ri==0) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
            set_font(p.add_run(str(val)),font_size,(ri==0 and header),color=(0,0,0))
    para(after=1)
    return t

def caption(text):
    p=para(align=WD_ALIGN_PARAGRAPH.CENTER,before=2,after=3,keep=True)
    set_font(p.add_run(text),8,color=(40,48,54))

def figure(path,width=3.28):
    p=para(align=WD_ALIGN_PARAGRAPH.CENTER,after=0,keep=True)
    p.add_run().add_picture(str(path),width=Inches(width))

p=para(after=1)
set_font(p.add_run("Abstract—"),9,True,italic=True)
set_font(p.add_run("This paper documents a reproducible UPI transaction fraud classification experiment and a small Flask demonstration built from a supplied CSV dataset. A random forest was trained on 200,000 records and evaluated on a stratified holdout of 50,000 records. On that split, accuracy was 92.98%, precision 93.09%, recall 92.85%, F1 score 92.97%, and ROC AUC 0.930. The confusion matrix contained 23,275 true negatives, 1,722 false positives, 1,787 false negatives, and 23,216 true positives. These results describe performance within this file and split. The observed fraud prevalence is 50.006%, and the dataset documentation does not establish how labels or the supplied risk score were generated. Since risk score and transaction amount contribute 88.0% of impurity-based feature importance, possible target leakage must be examined before any operational interpretation. The accompanying website reports the evaluation results and returns a model score for a submitted transaction; it is a prototype, not a payment control."),9,italic=True)
p=para(after=4)
set_font(p.add_run("Index Terms—"),9,True,italic=True)
set_font(p.add_run("UPI, fraud detection, random forest, imbalanced learning, confusion matrix, Flask."),9,italic=True)

section("I. INTRODUCTION")
body("UPI enables fast person-to-person and merchant payments, so a transaction screening model must distinguish suspicious activity without creating excessive false alerts. Machine learning can rank transactions using observed attributes, but a high score on a single dataset does not by itself establish that a model will work on live payments. The performance depends on the data source, label quality, class prevalence, feature availability, and evaluation design.")
body("This study reports the model already implemented in the project: a scikit-learn preprocessing pipeline followed by a class-weighted random forest, plus a Flask interface for displaying held-out metrics and scoring a form submission. The work contributes a transparent account of the dataset, split, feature exclusions, measured confusion matrix, and limits that affect interpretation. Prior UPI fraud-detection discussion motivates supervised classification and adaptive monitoring, while the present experiment is limited to a static CSV and does not test those operational capabilities [1].")
body("The main result is 92.98% accuracy on a stratified 20% holdout. Because the dataset is almost exactly balanced, this accuracy is not inflated by a dominant majority class in this particular split. However, its 50.006% fraud prevalence is not established as representative of real payment traffic. In addition, the `risk_score` field may encode information related to the target. We therefore treat the reported scores as an internal benchmark, not evidence of deployment readiness.")

section("II. RELATED WORK")
body("Random forests aggregate many decision trees and are commonly used when transaction data contain mixed numeric and categorical predictors [2]. The supplied chapter surveys machine-learning approaches to UPI fraud detection and discusses feature engineering, anomaly detection, and monitoring [1]. It provides context for this project but does not validate the dataset used here; the project CSV has no source or labeling protocol in its accompanying metadata.")
body("Classification accuracy alone can hide different costs for false alarms and missed fraud. ROC analysis compares the true-positive rate with the false-positive rate across thresholds, while precision and recall describe the positive class at a selected threshold [3]. For a fraud screen, the operating threshold should be chosen using the actual fraud prevalence and the relative costs of review and missed losses. This study reports several metrics and makes no claim that the default 0.5 threshold is optimal.")
body("The experiment uses scikit-learn’s pipeline and model-selection components [4]. A pipeline keeps imputation and encoding attached to the estimator, reducing the risk that preprocessing learns from the held-out rows. Stratification preserves the target proportions across the train and test partitions; it does not protect against repeated actors or time-related leakage.")

section("III. DATASET AND PREPROCESSING")
subhead("A. Dataset description")
body("The supplied file contains 250,000 rows and 20 columns, including the binary target `is_fraud`. There are 125,016 fraud labels and 124,984 legitimate labels, giving a file-level fraud rate of 50.006%. Transaction amounts range from INR 10 to INR 42,099, with a median of INR 629. The CSV includes transaction type, merchant category, status, sender and receiver age groups, state, banks, device and network type, hour, weekday, weekend flag, and risk score. The project does not document the collection method, label adjudication, or the meaning and derivation of the risk score.")
subhead("B. Feature handling")
body("The model excludes transaction ID, timestamp, sender UPI ID, receiver UPI ID, and the target column from its predictors. Excluding identifiers avoids direct memorization of transaction keys and personal payment addresses. The timestamp is not decomposed into additional calendar features in the present implementation; hour, weekday, and weekend fields already appear separately in the file. The remaining 15 predictors enter the pipeline.")
body("Numeric fields are imputed with their training-set median and standardized. Categorical fields are imputed with their training-set most frequent value and one-hot encoded; unseen categories are ignored at inference. The random forest itself does not require standardized magnitudes, so scaling is included as part of the existing project pipeline rather than as a claimed source of improved forest performance. All preprocessing is fitted inside each training run.")
subhead("C. Risk score as a potential leakage channel")
body("Risk score has the largest measured feature importance, followed by transaction amount. Together they account for about 88.0% of the forest’s impurity-based importance. Mean risk score is 1,163.06 for fraud rows and 295.39 for legitimate rows. This separation could reflect a useful signal, a score available before a decision, or a value constructed using the fraud label or related post-event evidence. The file does not identify which explanation applies. Until its provenance and timing are verified, model performance with this field should be treated as potentially optimistic.")

section("IV. MODEL AND EXPERIMENTAL DESIGN")
subhead("A. Classifier")
body("The project classifier is a random forest with 100 trees, maximum depth 20, minimum leaf size 2, and `balanced_subsample` class weighting. The random seed is 42 and four worker threads are used. A forest combines tree predictions and can represent nonlinear interactions between transaction attributes [2]. The specified class weighting adjusts each tree’s bootstrap sample for class counts; it does not change the underlying label prevalence in the test set.")
subhead("B. Train and test split")
body("The records are divided once using an 80/20 random stratified split with seed 42. Training uses 200,000 rows; testing uses 50,000 rows. The test set contains 24,997 legitimate and 25,003 fraud transactions. No test records are used to fit the imputers, encoders, or forest. This row-wise split is reproducible, but it does not hold out entire senders, receivers, merchants, or later time periods.")
subhead("C. Evaluation measures")
body("Accuracy is the share of correct predictions. Precision is the fraction of predicted fraud cases that are fraud; recall is the fraction of fraud cases detected. F1 is the harmonic mean of precision and recall. ROC AUC measures ranking across thresholds [3]. The confusion matrix records true negatives, false positives, false negatives, and true positives. The web application uses a probability threshold of 0.5 for its displayed class label. The probability output has not been calibrated against an independent population.")

section("V. RESULTS")
body("At the 0.5 threshold, the model correctly classified 46,491 of 50,000 test transactions. It identified 23,216 of 25,003 fraud-labelled transactions and correctly cleared 23,275 of 24,997 legitimate transactions. The corresponding false-positive rate is 6.89%, and the false-negative rate is 7.15%. Table I summarizes the main scores; Fig. 1 shows the underlying counts.")
caption("TABLE I\nHELD-OUT TEST PERFORMANCE")
table([["Measure","Value"],["Accuracy","0.9298"],["Precision","0.9309"],["Recall","0.9285"],["F1 score","0.9297"],["ROC AUC","0.9300"],["Test records","50,000"]],[1.75,1.2],True,8)
figure(figs["confusion_matrix"],3.18)
caption("Fig. 1. Confusion matrix for the stratified test set at a 0.5 probability threshold.")
body("The ROC curve in Fig. 2 summarizes ranking behaviour across thresholds. Its AUC is 0.930. This is consistent with useful separation on the held-out file, although it does not indicate the precision that would result at a much lower real-world fraud rate. For a fixed sensitivity and specificity, positive predictive value declines as prevalence declines. A deployment study would therefore need an appropriately sampled, temporally separated test set and should report precision-recall behaviour at the expected operating prevalence.")
figure(figs["roc_curve"],3.18)
caption("Fig. 2. ROC curve from the held-out test predictions; plotted points are interpolated at evenly spaced false-positive rates.")

section("VI. FEATURE IMPORTANCE AND INTERPRETATION")
body("Fig. 3 ranks the six largest aggregated input importances. Risk score contributes 48.5%, amount 39.5%, and merchant category 5.9%; all remaining fields contribute smaller shares. These values are impurity-based splits aggregated across one-hot columns. They measure how the fitted forest used features, not causal effects or independent predictive value. Correlated inputs can share or substitute importance, and high-cardinality variables can receive biased importance estimates.")
figure(figs["feature_importance"],3.18)
caption("Fig. 3. Largest aggregated random-forest feature importances.")
body("The concentration of importance in risk score and amount is a practical audit target. A controlled ablation should retrain the same pipeline without risk score and compare the holdout metrics. A second model should use a time-based split, and, where repeated participants exist, a grouped split should keep each participant on one side. Those analyses would help separate genuine early warning from target leakage and repeated-entity memorization. They were not part of the results reported here.")

section("VII. FLASK DEMONSTRATION")
body("The project exposes a local Flask page at the root route. It displays accuracy, precision, recall, ROC AUC, and the confusion matrix saved with the trained model. A transaction form submits feature values to `POST /api/predict`; the route applies the stored preprocessing and returns a fraud probability and a class label. `GET /api/metrics` returns the stored evaluation metrics. This makes the experiment inspectable without rerunning training for each page visit.")
body("The application is a demonstration interface. It has no authentication, audit log, rate limiting, model monitoring, or integration with a payment processor. It should remain local or behind appropriate controls. A prediction is not a decision to block a payment, and the current UI does not include human review or a calibrated risk policy.")

section("VIII. LIMITATIONS AND NEXT STEPS")
body("The results are limited by the dataset’s undocumented provenance and unusually high, nearly balanced fraud prevalence. The labels may not represent adjudicated financial fraud, and the risk score may carry target-derived information. The random row split can also place transactions from the same underlying actors in both partitions. Only one random seed and one model configuration were evaluated; no confidence interval, temporal test, ablation, baseline comparison, or external validation was performed. The reported probabilities have not been calibrated for production prevalence.")
body("Before operational use, the dataset owner should document label rules, sampling, data dates, and feature availability at decision time. The next evaluation should include a risk-score ablation, a time-based holdout, grouped validation where identities can be safely represented, and threshold selection tied to review capacity and fraud-loss costs. It should also measure precision at the expected base rate, calibration, subgroup error rates, and performance drift. These checks would clarify whether the observed discrimination survives a more realistic evaluation.")

section("IX. CONCLUSION")
body("A random-forest pipeline trained on the supplied 250,000-row dataset achieved 92.98% accuracy and 0.930 ROC AUC on a reproducible 50,000-row stratified test set. The confusion matrix and accompanying figures show similar error rates for the two nearly balanced classes. The result demonstrates that the project can train, evaluate, save, and serve a model through a Flask interface. It does not establish readiness for live UPI fraud decisions. Label provenance, the risk-score field, realistic prevalence, and time- or entity-separated validation remain unresolved and should be addressed before relying on the score.")

section("REFERENCES")
refs=[
"[1] S. Jagadeesan, K. S. Arjun, G. Dhanika, G. Karthikeyan, and K. Deepika, “UPI fraud detection using machine learning,” in Challenges in Information, Communication and Computing Technology, 2025, pp. 755–760, doi: 10.1201/9781003559085-130.",
"[2] L. Breiman, “Random forests,” Machine Learning, vol. 45, pp. 5–32, 2001, doi: 10.1023/A:1010933404324.",
"[3] T. Fawcett, “An introduction to ROC analysis,” Pattern Recognition Letters, vol. 27, no. 8, pp. 861–874, 2006, doi: 10.1016/j.patrec.2005.10.010.",
"[4] F. Pedregosa et al., “Scikit-learn: Machine learning in Python,” Journal of Machine Learning Research, vol. 12, no. 85, pp. 2825–2830, 2011."
]
for r in refs:
    p=para(first=0,after=2)
    set_font(p.add_run(r),8)

# Compact IEEE page number field; keep it unobtrusive.
footer=doc.sections[-1].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run(); set_font(run,8,color=GRAY)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)

# Set document metadata without inventing institutional affiliation.
doc.core_properties.title="UPI Fraud Detection with a Random Forest Model"
doc.core_properties.author="Mohit Singh"
doc.core_properties.subject="Research paper based on the supplied UPI fraud detection project"
doc.core_properties.keywords="UPI, fraud detection, random forest, machine learning"
doc.save(OUT)
print(OUT)
